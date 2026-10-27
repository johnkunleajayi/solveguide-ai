import asyncio
import copy
import io
from contextlib import asynccontextmanager
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.auth.dependencies import required_current_user
from app.auth.security import now
from app.config import Settings
from app.context_library.extraction import extract
from app.context_library.models import ContextDocument
from app.context_library.service import ContextService
from app.context_library.storage import LocalDocumentStorage, new_storage_key
from app.database import PlatformError
from app.main import create_app

ORIGIN = {"Origin": "http://tauri.localhost"}
BASE = "/api/v1/context/documents"


class Session:
    def __init__(self, rows):
        self.rows = rows

    def matches(self, statement):
        params = statement.compile().params
        # Real SQLAlchemy statements, service-boundary fake (not a PostgreSQL substitute).
        assert "owner_id_1" in params, "Every document query must enforce ownership"
        rows = [row for row in self.rows if row.owner_id == params["owner_id_1"]]
        if "id_1" in params and "created_at_1" not in params:
            rows = [row for row in rows if row.id == params["id_1"]]
        if "created_at_1" in params:
            rows = [
                row
                for row in rows
                if (row.created_at, row.id) < (params["created_at_1"], params["id_1"])
            ]
        return sorted(rows, key=lambda r: (r.created_at, r.id), reverse=True)

    async def scalar(self, statement):
        return next(iter(self.matches(statement)), None)

    async def scalars(self, statement):
        return SimpleNamespace(all=lambda: self.matches(statement)[:51])

    def add(self, row):
        self.rows.append(row)

    async def flush(self):
        for row in self.rows:
            row.id = row.id or uuid4()
            row.created_at = row.created_at or now()
            row.updated_at = now()

    async def delete(self, row):
        self.rows.remove(row)


class Database:
    def __init__(self):
        self.rows, self.fail_commit = [], False

    @asynccontextmanager
    async def transaction(self):
        snapshot = copy.deepcopy(self.rows)
        try:
            yield Session(self.rows)
            if self.fail_commit:
                self.fail_commit = False
                raise PlatformError("database_unavailable", "Unavailable")
        except BaseException:
            self.rows = snapshot
            raise


@pytest.fixture
def library(tmp_path):
    settings = Settings(_env_file=None, context_storage_directory=tmp_path / "documents")
    app = create_app(settings)
    db = Database()
    storage = LocalDocumentStorage(settings.context_storage_directory)
    app.state.context_service = ContextService(db, storage, settings)
    user = SimpleNamespace(id=uuid4())
    app.dependency_overrides[required_current_user] = lambda: user
    app.state.vision_service = Mock()
    app.state.answer_service = Mock()
    with TestClient(app, base_url="http://127.0.0.1") as client:
        yield SimpleNamespace(
            client=client, app=app, db=db, storage=storage, user=user, settings=settings
        )
    assert app.state.vision_service.mock_calls == []
    assert app.state.answer_service.mock_calls == []


def upload(lib, name="notes.txt", content=b"A reference note.", media="text/plain"):
    return lib.client.post(
        BASE + "/upload",
        headers={**ORIGIN, "Content-Type": media, "X-Document-Filename": name},
        content=content,
    )


def pdf(text=True, pages=1):
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=300, height=300)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    )
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 20 20 Td (A practice reference.) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_all_library_routes_require_authentication():
    with TestClient(create_app(Settings(_env_file=None)), base_url="http://127.0.0.1") as client:
        for method, path in [
            ("GET", BASE),
            ("GET", BASE + "/limits"),
            ("POST", BASE + "/upload"),
            ("POST", BASE + "/text"),
            ("GET", BASE + "/" + str(uuid4())),
            ("PATCH", BASE + "/" + str(uuid4())),
            ("DELETE", BASE + "/" + str(uuid4())),
        ]:
            response = client.request(method, path, headers=ORIGIN)
            assert response.status_code == 401


def test_txt_upload_preview_and_no_internal_storage_fields(library):
    response = upload(library)
    assert response.status_code == 201
    doc = response.json()
    assert doc["extracted_text"] == "A reference note."
    assert doc["extraction_status"] == "ready"
    assert "storage_key" not in doc and "owner_id" not in doc
    assert library.client.get(BASE).json()["documents"][0]["id"] == doc["id"]
    assert "extracted_text" not in library.client.get(BASE).json()["documents"][0]
    assert len(list(library.storage.root.iterdir())) == 1


def test_utf16_txt(library):
    assert (
        upload(library, content="Résumé notes".encode("utf-16")).json()["extracted_text"]
        == "Résumé notes"
    )


def test_pdf_pages_and_offsets(library):
    response = upload(library, "guide.pdf", pdf(pages=2), "application/pdf")
    assert response.status_code == 201
    doc = response.json()
    assert doc["extraction_status"] == "ready"
    assert [s["number"] for s in doc["extraction_structure"]] == [1, 2]
    for segment in doc["extraction_structure"]:
        assert doc["extracted_text"][segment["start"] : segment["end"]] == "A practice reference."


def test_image_only_pdf_is_not_successful_extraction(library):
    doc = upload(library, "scanned.pdf", pdf(text=False), "application/pdf").json()
    assert doc["extraction_status"] == "insufficient"
    assert doc["extraction_error"] == "no_readable_text"


def test_docx_paragraph_table_reading_order(library):
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Left", "Right"
    document.add_paragraph("After")
    data = io.BytesIO()
    document.save(data)
    doc = upload(
        library,
        "guide.docx",
        data.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ).json()
    assert doc["extracted_text"] == "Before\n\nLeft\tRight\n\nAfter"
    assert [s["kind"] for s in doc["extraction_structure"]] == ["paragraph", "table", "paragraph"]


def test_paste_has_no_binary_and_normalizes_text(library):
    doc = library.client.post(
        BASE + "/text", headers=ORIGIN, json={"title": " Notes ", "text": "First\r\nSecond"}
    ).json()
    assert doc["title"] == "Notes" and doc["extracted_text"] == "First\nSecond"
    assert doc["source_type"] == "pasted_text" and doc["original_filename"] is None
    assert not library.storage.root.exists()


@pytest.mark.parametrize(
    "name,media,data,status",
    [
        ("a.exe", "application/octet-stream", b"bad", 415),
        ("a.txt", "application/pdf", b"bad", 415),
        ("..%2Fsecret.txt", "text/plain", b"text", 422),
        ("a.txt", "text/plain", b"\x00\x01binary", 422),
        ("a.txt", "text/plain", b"\xffbad", 422),
        ("a.txt", "text/plain", b"", 422),
    ],
)
def test_invalid_uploads_rejected(library, name, media, data, status):
    assert upload(library, name, data, media).status_code == status
    assert not library.db.rows


@pytest.mark.parametrize(
    "filename,media",
    [
        ("bad.pdf", "application/pdf"),
        ("bad.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ],
)
def test_malformed_parser_failure_is_sanitized(library, filename, media):
    import zipfile

    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("word/document.xml", "private document parser detail")
    content = (
        b"%PDF-1.7\nprivate document parser detail"
        if filename.endswith(".pdf")
        else data.getvalue()
    )
    doc = upload(library, filename, content, media).json()
    assert doc["extraction_status"] == "failed" and doc["extracted_text"] == ""
    assert doc["extraction_error"] == "extraction_failed"


def test_upload_extraction_and_page_limits(library):
    library.settings.context_upload_bytes = 1024
    assert upload(library, content=b"a" * 1025).status_code == 413
    library.settings.context_upload_bytes = 10 * 1024 * 1024
    library.settings.context_text_characters = 1000
    assert upload(library, content=b"a" * 1001).status_code == 413
    library.settings.context_max_pages = 1
    assert upload(library, "guide.pdf", pdf(pages=2), "application/pdf").status_code == 413


def test_cross_user_list_detail_update_and_delete_fail_closed(library):
    doc = upload(library).json()
    library.user.id = uuid4()
    assert library.client.get(BASE).json()["documents"] == []
    path = BASE + "/" + doc["id"]
    for method in ["GET", "PATCH", "DELETE"]:
        result = library.client.request(
            method, path, headers=ORIGIN, json={"title": "Hijacked"} if method == "PATCH" else None
        )
        assert result.status_code == 404
    assert library.db.rows[0].title == "notes.txt"
    assert len(list(library.storage.root.iterdir())) == 1


def test_rename_and_delete_remove_all_content(library):
    doc = upload(library).json()
    path = BASE + "/" + doc["id"]
    assert (
        library.client.patch(path, headers=ORIGIN, json={"title": "New title"}).json()["title"]
        == "New title"
    )
    assert library.client.delete(path, headers=ORIGIN).status_code == 204
    assert library.db.rows == [] and list(library.storage.root.iterdir()) == []


def test_storage_delete_failure_keeps_retryable_tombstone(library, monkeypatch):
    doc = upload(library).json()
    original = library.storage.delete

    def fail(_):
        raise PlatformError("storage_unavailable", "Deletion is pending.")

    monkeypatch.setattr(library.storage, "delete", fail)
    path = BASE + "/" + doc["id"]
    assert library.client.delete(path, headers=ORIGIN).status_code == 503
    row = library.db.rows[0]
    assert row.extraction_status == "deleting" and row.extracted_text == ""
    assert library.client.patch(path, headers=ORIGIN, json={"title": "No"}).status_code == 409
    monkeypatch.setattr(library.storage, "delete", original)
    assert library.client.delete(path, headers=ORIGIN).status_code == 204
    assert library.db.rows == [] and list(library.storage.root.iterdir()) == []


def test_failed_database_commit_compensates_binary(library):
    library.db.fail_commit = True
    assert upload(library).status_code == 503
    assert library.db.rows == [] and list(library.storage.root.iterdir()) == []


def test_delete_database_failure_after_binary_removal_is_retryable(library, monkeypatch):
    doc = upload(library).json()
    original = library.storage.delete

    def remove_then_fail_commit(key):
        original(key)
        library.db.fail_commit = True

    monkeypatch.setattr(library.storage, "delete", remove_then_fail_commit)
    path = BASE + "/" + doc["id"]
    assert library.client.delete(path, headers=ORIGIN).status_code == 503
    assert library.db.rows[0].extraction_status == "deleting"
    monkeypatch.setattr(library.storage, "delete", original)
    assert library.client.delete(path, headers=ORIGIN).status_code == 204


def test_origins_and_title_validation(library):
    assert (
        library.client.post(
            BASE + "/text",
            headers={"Origin": "https://other.test"},
            json={"title": "Title", "text": "Notes"},
        ).status_code
        == 403
    )
    assert (
        library.client.post(
            BASE + "/text", headers=ORIGIN, json={"title": " ", "text": "Notes"}
        ).status_code
        == 422
    )
    assert (
        library.client.post(
            BASE + "/text", headers=ORIGIN, json={"title": "Title", "text": "   "}
        ).status_code
        == 422
    )


def test_storage_roundtrip_missing_and_traversal(tmp_path):
    storage = LocalDocumentStorage(tmp_path / "private")
    key = new_storage_key()
    storage.write(key, b"notes")
    assert storage.read(key) == b"notes"
    storage.delete(key)
    storage.delete(key)
    with pytest.raises(PlatformError) as error:
        storage.read(key)
    assert error.value.code == "document_not_found"
    for invalid in ["../secret", "..\\secret", "C:\\secret", "/absolute", "a" * 31, "G" * 32]:
        with pytest.raises(PlatformError):
            storage.write(invalid, b"bad")


def test_failed_storage_write_does_not_create_database_row(library, monkeypatch):
    def fail(*_):
        raise PlatformError("storage_unavailable", "Unavailable")

    monkeypatch.setattr(library.storage, "write", fail)
    assert upload(library).status_code == 503
    assert library.db.rows == []


def test_partial_storage_write_is_cleaned_up(tmp_path, monkeypatch):
    import os

    storage = LocalDocumentStorage(tmp_path / "private")

    def fail(_):
        raise OSError("disk failure")

    monkeypatch.setattr(os, "fsync", fail)
    with pytest.raises(PlatformError):
        storage.write(new_storage_key(), b"partial file")
    assert list(storage.root.iterdir()) == []


@pytest.mark.parametrize("failure", ["session_expired", "session_revoked"])
def test_expired_or_revoked_session_cannot_read_library(library, failure):
    from datetime import timedelta

    from test_auth import MemoryDatabase

    from app.auth.models import AuthSession, User
    from app.auth.security import digest, secret

    token = secret()
    db = MemoryDatabase()
    user = User(
        id=library.user.id,
        google_subject="library-test",
        email="test@example.com",
        email_verified=True,
        display_name="Test",
    )
    session = AuthSession(
        id=uuid4(),
        user_id=user.id,
        token_hash=digest(token),
        created_at=now() - timedelta(hours=2),
        expires_at=now() - timedelta(hours=1)
        if failure == "session_expired"
        else now() + timedelta(hours=1),
        last_seen_at=now(),
        revoked_at=now() if failure == "session_revoked" else None,
    )
    db.rows[User].append(user)
    db.rows[AuthSession].append(session)
    library.app.dependency_overrides.clear()
    library.app.state.database = db
    response = library.client.get(BASE, headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 401 and response.json()["error"]["code"] == failure
    assert library.db.rows == []


def test_plain_text_cannot_execute_markup(library):
    doc = library.client.post(
        BASE + "/text",
        headers=ORIGIN,
        json={"title": "Reference", "text": "<script>alert(1)</script>"},
    ).json()
    assert doc["extracted_text"] == "<script>alert(1)</script>"


def test_page_cursor_is_bounded_and_owner_scoped(library):
    async def seed():
        async with library.db.transaction() as db:
            for n in range(52):
                db.add(
                    ContextDocument(
                        owner_id=library.user.id,
                        title=str(n),
                        original_filename=None,
                        document_type="txt",
                        source_type="pasted_text",
                        storage_key=None,
                        byte_size=1,
                        extraction_status="ready",
                        extraction_error=None,
                        extracted_text="a",
                        extraction_structure=[],
                        character_count=1,
                    )
                )
            await db.flush()

    asyncio.run(seed())
    first = library.client.get(BASE).json()
    second = library.client.get(BASE, params={"cursor": first["next_cursor"]}).json()
    assert len(first["documents"]) == 50 and len(second["documents"]) == 2
    assert library.client.get(BASE, params={"cursor": "invalid"}).status_code == 422


def test_zip_expansion_limit_and_no_docx_xml_entities(library):
    import zipfile

    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr(
            "word/document.xml", '<!DOCTYPE evil [<!ENTITY x SYSTEM="file:///secret">]><a/>'
        )
    result = extract(data.getvalue(), "docx", library.settings)
    assert result.status == "failed" and result.error == "extraction_failed"
    library.settings.context_expanded_bytes = 10
    with pytest.raises(PlatformError) as error:
        extract(data.getvalue(), "docx", library.settings)
    assert error.value.code == "extraction_limit"
